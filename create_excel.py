
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

folder = Path(__file__).resolve().parent / "papers"
input_file = folder / "catalog.csv"
output_file = folder / "catalog.xlsx"

with input_file.open("r", encoding="utf-8-sig", newline="") as f:
    rows = list(csv.reader(f))

if not rows:
    raise ValueError("CSV file is empty")

headers = rows[0]
if any(len(row) != len(headers) for row in rows):
    raise ValueError("Some rows have different column counts")

wb = Workbook()
ws = wb.active
ws.title = "Research Papers"

ws.append(headers)

for row in rows[1:]:
    ws.append(row)

for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="17365D")
    cell.alignment = Alignment(wrap_text=True)

for col in ws.columns:
    letter = get_column_letter(col[0].column)
    max_length = max(
        len(str(cell.value or "")) for cell in col
    )
    ws.column_dimensions[letter].width = min(
        max(max_length + 2, 15), 55
    )

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top", wrap_text=True
        )
    ws.row_dimensions[row[0].row].height = 45

ws.freeze_panes = "A2"

if ws.max_row > 1:
    table = Table(
        displayName="ResearchCatalog",
        ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True
    )
    ws.add_table(table)

wb.save(output_file)

print("Excel file created:", output_file)
print("Research records:", len(rows) - 1)
print("Columns:", len(headers))
